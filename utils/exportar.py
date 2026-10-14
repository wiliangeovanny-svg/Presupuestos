import io
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import LineChart, BarChart, Reference

FONT_NAME = "Arial"
HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
SUBHEADER_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
THIN_BORDER = Border(*[Side(style="thin", color="B7B7B7")] * 4)

TIPO_LABEL = {"material": "MATERIALES", "mano_obra": "MANO DE OBRA", "equipo": "EQUIPOS Y HERRAMIENTA"}


def _header_cell(ws, cell, text, fill=HEADER_FILL, color="FFFFFF", size=11, bold=True):
    c = ws[cell]
    c.value = text
    c.font = Font(name=FONT_NAME, bold=bold, color=color, size=size)
    c.fill = fill
    c.alignment = Alignment(horizontal="center", vertical="center")
    return c


def _sanitizar_nombre_hoja(nombre: str, usados: set) -> str:
    """Excel limita nombres de hoja a 31 caracteres y prohíbe : \\ / ? * [ ].
    Además evita duplicados agregando un sufijo numérico si hace falta."""
    for ch in r':\/?*[]':
        nombre = nombre.replace(ch, "-")
    nombre = nombre.strip()[:31] or "APU"
    base = nombre
    i = 2
    while nombre in usados:
        sufijo = f" ({i})"
        nombre = base[: 31 - len(sufijo)] + sufijo
        i += 1
    usados.add(nombre)
    return nombre


def _escribir_apu_en_hoja(ws, rubro: dict, resultado: dict):
    """Escribe el detalle completo de un APU en una hoja ya creada."""
    for col, width in zip("ABCDE", [38, 12, 12, 14, 14]):
        ws.column_dimensions[col].width = width

    ws.merge_cells("A1:E1")
    _header_cell(ws, "A1", "ANÁLISIS DE PRECIO UNITARIO (APU)", size=13)

    ws["A3"] = "Rubro:"
    ws["A3"].font = Font(name=FONT_NAME, bold=True)
    ws["B3"] = rubro.get("descripcion", "")
    ws["D3"] = "Unidad:"
    ws["D3"].font = Font(name=FONT_NAME, bold=True)
    ws["E3"] = rubro.get("unidad", "")

    row = 5
    grupos = {
        "material": [d for d in resultado["detalle"] if d["tipo"] == "material"],
        "mano_obra": [d for d in resultado["detalle"] if d["tipo"] == "mano_obra"],
        "equipo": [d for d in resultado["detalle"] if d["tipo"] == "equipo"],
    }
    subtotales = {
        "material": resultado["subtotal_materiales"],
        "mano_obra": resultado["subtotal_mano_obra"],
        "equipo": resultado["subtotal_equipo"],
    }

    for tipo, items in grupos.items():
        ws.merge_cells(f"A{row}:E{row}")
        _header_cell(ws, f"A{row}", TIPO_LABEL[tipo], fill=SUBHEADER_FILL, color="1F4E78")
        row += 1

        headers = ["Descripción", "Unidad", "Cantidad", "Precio Unit.", "Costo"]
        for i, h in enumerate(headers):
            cell = ws.cell(row=row, column=i + 1, value=h)
            cell.font = Font(name=FONT_NAME, bold=True, size=10)
            cell.border = THIN_BORDER
        row += 1

        for it in items:
            ws.cell(row=row, column=1, value=it.get("descripcion", "")).border = THIN_BORDER
            ws.cell(row=row, column=2, value=it.get("unidad", "")).border = THIN_BORDER
            ws.cell(row=row, column=3, value=it.get("cantidad", 0)).border = THIN_BORDER
            ws.cell(row=row, column=4, value=round(it.get("precio", 0), 4)).border = THIN_BORDER
            ws.cell(row=row, column=5, value=round(it.get("costo", 0), 4)).border = THIN_BORDER
            row += 1

        ws.cell(row=row, column=4, value=f"Subtotal {TIPO_LABEL[tipo]}").font = Font(name=FONT_NAME, bold=True, size=9)
        ws.cell(row=row, column=5, value=round(subtotales[tipo], 4)).font = Font(name=FONT_NAME, bold=True)
        row += 2

    ws.cell(row=row, column=4, value="COSTO DIRECTO").font = Font(name=FONT_NAME, bold=True)
    ws.cell(row=row, column=5, value=round(resultado["costo_directo"], 4)).font = Font(name=FONT_NAME, bold=True)
    row += 1

    ws.cell(row=row, column=4, value=f"Indirectos y Utilidad ({resultado['indirectos_utilidad_pct']}%)")
    ws.cell(row=row, column=5, value=round(resultado["monto_indirectos_utilidad"], 4))
    row += 1

    ws.cell(row=row, column=4, value="PRECIO UNITARIO").font = Font(name=FONT_NAME, bold=True, size=11)
    ws.cell(row=row, column=4).fill = SUBHEADER_FILL
    pu_cell = ws.cell(row=row, column=5, value=round(resultado["precio_unitario"], 4))
    pu_cell.font = Font(name=FONT_NAME, bold=True, size=11)
    pu_cell.fill = SUBHEADER_FILL


def exportar_apu_excel(rubro: dict, resultado: dict) -> bytes:
    """Genera un Excel independiente con el formato clásico de un solo APU."""
    wb = Workbook()
    ws = wb.active
    ws.title = "APU"
    _escribir_apu_en_hoja(ws, rubro, resultado)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def exportar_presupuesto_excel(
    proyecto: dict,
    resultado: dict,
    apus: list[dict] | None = None,
    cronograma: tuple[list[str], dict] | None = None,
) -> bytes:
    """Genera un Excel con el presupuesto completo del proyecto.

    apus: lista opcional de {"rubro": {...}, "resultado": {...}} — uno por cada
    rubro distinto usado en el presupuesto. Si se pasa, se agrega una hoja
    adicional por cada rubro con el detalle completo de su APU, para que el
    archivo quede autocontenido (presupuesto + respaldo de cada precio).

    cronograma: opcional, tupla (etiquetas, resultado_cronograma) tal como la
    devuelve utils.negocio.calcular_cronograma_proyecto. Si se pasa, agrega
    una hoja "Cronograma Valorado" con la matriz y la curva S. Es opcional a
    propósito: solo aplica a los proyectos que activaron el cronograma.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Presupuesto"

    for col, width in zip("ABCDEF", [8, 42, 10, 12, 14, 16]):
        ws.column_dimensions[col].width = width

    ws.merge_cells("A1:F1")
    _header_cell(ws, "A1", f"PRESUPUESTO — {proyecto.get('nombre', '')}", size=13)
    if proyecto.get("ubicacion"):
        ws.merge_cells("A2:F2")
        ws["A2"] = f"Ubicación: {proyecto['ubicacion']}"
        ws["A2"].font = Font(name=FONT_NAME, italic=True)

    row = 4
    headers = ["N°", "Descripción", "Unidad", "Cantidad", "Precio Unit.", "Subtotal"]
    for i, h in enumerate(headers):
        cell = ws.cell(row=row, column=i + 1, value=h)
        cell.font = Font(name=FONT_NAME, bold=True, color="FFFFFF", size=10)
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER
        cell.alignment = Alignment(horizontal="center")
    row += 1

    for i, it in enumerate(resultado["filas"], start=1):
        ws.cell(row=row, column=1, value=i).border = THIN_BORDER
        ws.cell(row=row, column=2, value=it.get("rubro_descripcion", "")).border = THIN_BORDER
        ws.cell(row=row, column=3, value=it.get("unidad", "")).border = THIN_BORDER
        ws.cell(row=row, column=4, value=it.get("cantidad_obra", 0)).border = THIN_BORDER
        ws.cell(row=row, column=5, value=round(it.get("precio_unitario", 0), 4)).border = THIN_BORDER
        ws.cell(row=row, column=6, value=round(it.get("subtotal", 0), 4)).border = THIN_BORDER
        row += 1

    row += 1
    ws.cell(row=row, column=5, value="TOTAL PRESUPUESTO").font = Font(name=FONT_NAME, bold=True, size=11)
    total_cell = ws.cell(row=row, column=6, value=round(resultado["total"], 4))
    total_cell.font = Font(name=FONT_NAME, bold=True, size=11)
    total_cell.fill = SUBHEADER_FILL

    nombres_usados = {ws.title}

    # Hoja del cronograma valorado (opcional, solo si el proyecto lo tiene activo).
    if cronograma:
        etiquetas_cron, resultado_cron = cronograma
        nombre_hoja_cron = _sanitizar_nombre_hoja("Cronograma Valorado", nombres_usados)
        ws_cron = wb.create_sheet(title=nombre_hoja_cron)
        _escribir_cronograma_en_hoja(ws_cron, proyecto, etiquetas_cron, resultado_cron)

    # Una hoja adicional por cada rubro distinto, con el detalle de su APU.
    if apus:
        for item in apus:
            rubro = item["rubro"]
            resultado_apu = item["resultado"]
            nombre_base = f"{rubro.get('codigo') or ''} {rubro['descripcion']}".strip()
            nombre_hoja = _sanitizar_nombre_hoja(nombre_base, nombres_usados)
            ws_apu = wb.create_sheet(title=nombre_hoja)
            _escribir_apu_en_hoja(ws_apu, rubro, resultado_apu)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _escribir_cronograma_en_hoja(ws, proyecto: dict, etiquetas: list[str], resultado: dict):
    """Escribe el cronograma valorado (matriz + curva S) en una hoja ya creada."""
    num_periodos = len(etiquetas)
    primera_col_periodo = 3  # A: rubro, B: subtotal, C en adelante: periodos

    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 14
    for i in range(num_periodos):
        ws.column_dimensions[get_column_letter(primera_col_periodo + i)].width = 13

    titulo_col_final = get_column_letter(primera_col_periodo + num_periodos - 1)
    ws.merge_cells(f"A1:{titulo_col_final}1")
    _header_cell(ws, "A1", f"CRONOGRAMA VALORADO — {proyecto.get('nombre', '')}", size=13)

    row = 3
    ws.cell(row=row, column=1, value="Rubro").font = Font(name=FONT_NAME, bold=True, color="FFFFFF")
    ws.cell(row=row, column=2, value="Valor Total").font = Font(name=FONT_NAME, bold=True, color="FFFFFF")
    for i, etiqueta in enumerate(etiquetas):
        ws.cell(row=row, column=primera_col_periodo + i, value=etiqueta).font = Font(
            name=FONT_NAME, bold=True, color="FFFFFF", size=9
        )
    for c in range(1, primera_col_periodo + num_periodos):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    row += 1

    fila_inicio_datos = row
    ws.freeze_panes = f"C{fila_inicio_datos}"  # inmoviliza col. Rubro/Valor Total y fila de encabezados
    for fila in resultado["filas"]:
        ws.cell(row=row, column=1, value=fila["rubro_descripcion"]).border = THIN_BORDER
        ws.cell(row=row, column=2, value=round(fila["subtotal"], 2)).border = THIN_BORDER
        for i, valor in enumerate(fila["valores"]):
            c = ws.cell(row=row, column=primera_col_periodo + i, value=round(valor, 2) if valor else None)
            c.border = THIN_BORDER
        row += 1
    fila_fin_datos = row - 1

    # Fila de totales por periodo
    ws.cell(row=row, column=1, value="TOTAL PERIODO").font = Font(name=FONT_NAME, bold=True)
    ws.cell(row=row, column=1).fill = SUBHEADER_FILL
    ws.cell(row=row, column=2, value=round(resultado["total_general"], 2)).font = Font(name=FONT_NAME, bold=True)
    ws.cell(row=row, column=2).fill = SUBHEADER_FILL
    for i, valor in enumerate(resultado["totales_periodo"]):
        c = ws.cell(row=row, column=primera_col_periodo + i, value=round(valor, 2))
        c.font = Font(name=FONT_NAME, bold=True)
        c.fill = SUBHEADER_FILL
    fila_totales = row
    row += 1

    # Fila de acumulado (curva S)
    ws.cell(row=row, column=1, value="ACUMULADO")
    for i, valor in enumerate(resultado["acumulado"]):
        ws.cell(row=row, column=primera_col_periodo + i, value=round(valor, 2))
    fila_acumulado = row
    row += 1

    # Fila de % acumulado
    ws.cell(row=row, column=1, value="% ACUMULADO")
    for i, valor in enumerate(resultado["pct_acumulado"]):
        ws.cell(row=row, column=primera_col_periodo + i, value=round(valor, 1))
    row += 1

    # --- Gráfico: curva S (valor acumulado) ---
    if num_periodos > 1:
        line = LineChart()
        line.title = "Curva S — Valor Acumulado"
        line.y_axis.title = "Valor acumulado"
        line.x_axis.title = "Periodo"
        line.style = 10
        datos_acum = Reference(ws, min_col=primera_col_periodo, max_col=primera_col_periodo + num_periodos - 1,
                                min_row=fila_acumulado, max_row=fila_acumulado)
        cats = Reference(ws, min_col=primera_col_periodo, max_col=primera_col_periodo + num_periodos - 1,
                          min_row=3, max_row=3)
        line.add_data(datos_acum, titles_from_data=False, from_rows=True)
        line.set_categories(cats)
        line.series[0].tx = None
        line.series[0].smooth = False
        ws.add_chart(line, f"A{fila_fin_datos + 8}")

        bar = BarChart()
        bar.title = "Valor Planificado por Periodo"
        bar.y_axis.title = "Valor"
        bar.x_axis.title = "Periodo"
        datos_periodo = Reference(ws, min_col=primera_col_periodo, max_col=primera_col_periodo + num_periodos - 1,
                                   min_row=fila_totales, max_row=fila_totales)
        bar.add_data(datos_periodo, titles_from_data=False, from_rows=True)
        bar.set_categories(cats)
        bar.series[0].tx = None
        ws.add_chart(bar, f"A{fila_fin_datos + 24}")


def exportar_cronograma_excel(proyecto: dict, etiquetas: list[str], resultado: dict) -> bytes:
    """Genera un Excel independiente con el cronograma valorado (matriz + curva S)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Cronograma Valorado"
    _escribir_cronograma_en_hoja(ws, proyecto, etiquetas, resultado)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
